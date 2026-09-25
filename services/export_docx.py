from pathlib import Path

from docx import Document


def _safe_stem(stem):
    return "".join(
        c if c not in '\\/:*?"<>|' else "_"
        for c in stem
    )


def export_docx(stem, text, summary, out_dir):
    name = f"{_safe_stem(stem)}_analysis.docx"

    output_path = Path(out_dir) / name

    doc = Document()

    doc.add_heading(
        "강의 분석 결과",
        level=1
    )

    doc.add_paragraph(
        summary.get("overview", "")
    )

    doc.add_heading(
        "핵심 내용",
        level=2
    )

    for item in summary.get("key_points", []):
        doc.add_paragraph(
            item,
            style="List Bullet"
        )

    doc.add_heading(
        "키워드",
        level=2
    )

    doc.add_paragraph(
        ", ".join(
            summary.get("keywords", [])
        )
    )

    doc.add_heading(
        "전사 원문",
        level=2
    )

    doc.add_paragraph(text)

    doc.save(output_path)

    return name