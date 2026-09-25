from pathlib import Path
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document


def _safe_stem(stem):
    return "".join(c if c not in '\\/:*?"<>|' else "_" for c in stem)


def export_txt(stem, text, out_dir):
    name = f"{_safe_stem(stem)}_transcript.txt"
    (Path(out_dir) / name).write_text(text, encoding="utf-8")
    return name


def _srt_time(seconds):
    ms = int(round(seconds * 1000))
    h, ms = divmod(ms, 3600000)
    m, ms = divmod(ms, 60000)
    s, ms = divmod(ms, 1000)
    return f"{h:02}:{m:02}:{s:02},{ms:03}"


def export_srt(stem, segments, out_dir):
    name = f"{_safe_stem(stem)}_transcript.srt"
    lines = []
    for i, seg in enumerate(segments, 1):
        lines += [
            str(i),
            f"{_srt_time(seg['start'])} --> {_srt_time(seg['end'])}",
            seg["text"],
            "",
        ]
    (Path(out_dir) / name).write_text("\n".join(lines), encoding="utf-8-sig")
    return name


def export_docx(stem, text, summary, out_dir):
    name = f"{_safe_stem(stem)}_analysis.docx"
    doc = Document()
    doc.add_heading("강의 분석 결과", level=1)
    doc.add_paragraph(summary.get("overview", ""))
    doc.add_heading("핵심 내용", level=2)
    for item in summary.get("key_points", []):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("키워드", level=2)
    doc.add_paragraph(", ".join(summary.get("keywords", [])))
    doc.add_heading("전사 원문", level=2)
    doc.add_paragraph(text)
    doc.save(Path(out_dir) / name)
    return name


def export_pdf(stem, text, summary, out_dir):
    name = f"{_safe_stem(stem)}_analysis.pdf"
    path = Path(out_dir) / name

    font_candidates = [
        Path("C:/Windows/Fonts/malgun.ttf"),
        Path("C:/Windows/Fonts/NanumGothic.ttf"),
    ]
    font_path = next((p for p in font_candidates if p.exists()), None)
    if not font_path:
        raise RuntimeError("한글 PDF용 폰트를 찾지 못했습니다. C:/Windows/Fonts/malgun.ttf를 확인하세요.")

    pdfmetrics.registerFont(TTFont("Korean", str(font_path)))
    styles = getSampleStyleSheet()
    styles["BodyText"].fontName = "Korean"
    styles["Title"].fontName = "Korean"
    styles["Heading2"].fontName = "Korean"

    doc = SimpleDocTemplate(str(path), pagesize=A4)
    story = [
        Paragraph("강의 분석 결과", styles["Title"]),
        Spacer(1, 12),
        Paragraph(summary.get("overview", ""), styles["BodyText"]),
        Spacer(1, 12),
        Paragraph("핵심 내용", styles["Heading2"]),
    ]
    for item in summary.get("key_points", []):
        story.append(Paragraph(f"• {item}", styles["BodyText"]))
    story += [
        Spacer(1, 12),
        Paragraph("전사 원문", styles["Heading2"]),
        Paragraph(text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>"), styles["BodyText"]),
    ]
    doc.build(story)
    return name
